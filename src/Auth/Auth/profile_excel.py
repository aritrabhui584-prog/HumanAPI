import shutil
import time
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment


EXCEL_FILE = Path(__file__).resolve().parent / "users.xlsx"

def _setup_worksheet_headers(sheet):
    sheet.title = "Users"
    headers = [
        "User ID",
        "Name",
        "Email",
        "Phone Number",
        "Email Verified",
        "Phone Verified",
        "Account Active",
        "Failed Login Attempts",
        "Account Locked Until",
        "Role",
        "Registration Date",
        "Last Updated"
    ]

    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center")

    widths = {
        "A": 12, "B": 20, "C": 30, "D": 18,
        "E": 18, "F": 18, "G": 18, "H": 25,
        "I": 25, "J": 15, "K": 22, "L": 22
    }

    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

def create_users_excel():
    workbook = Workbook()
    sheet = workbook.active
    _setup_worksheet_headers(sheet)
    
    try:
        workbook.save(EXCEL_FILE)
        print(f"Users Excel file created successfully at {EXCEL_FILE}")
    except Exception as e:
        print(f"Failed to save Excel file: {e}")
    finally:
        workbook.close()

def _backup_corrupted_file():
    if EXCEL_FILE.exists():
        timestamp = int(time.time())
        backup_path = EXCEL_FILE.with_name(f"users_corrupted_backup_{timestamp}.xlsx")
        try:
            shutil.move(str(EXCEL_FILE), str(backup_path))
            print(f"Corrupted Excel file backed up to {backup_path}")
        except Exception as e:
            print(f"Failed to backup corrupted file: {e}")

def add_user_to_excel(
    user_id,
    name,
    email,
    phone_number,
    email_verified,
    phone_verified,
    account_active,
    failed_login_attempts,
    account_locked_until,
    role,
    registration_date,
    last_updated
):
    if not EXCEL_FILE.exists():
        create_users_excel()

    workbook = None
    try:
        workbook = load_workbook(EXCEL_FILE)
    except Exception as e:
        print(f"Error loading Excel file, possibly corrupted: {e}")
        _backup_corrupted_file()
        create_users_excel()
        workbook = load_workbook(EXCEL_FILE)

    try:
        if "Users" not in workbook.sheetnames:
            sheet = workbook.create_sheet("Users")
            _setup_worksheet_headers(sheet)
        else:
            sheet = workbook["Users"]

        # Ensure datetime/None types are converted to strings if needed, 
        # or just append as openpyxl supports basic types
        sheet.append([
            str(user_id) if user_id is not None else "",
            name,
            email,
            phone_number,
            bool(email_verified),
            bool(phone_verified),
            bool(account_active),
            int(failed_login_attempts) if failed_login_attempts is not None else 0,
            str(account_locked_until) if account_locked_until is not None else "",
            str(role) if role is not None else "",
            str(registration_date) if registration_date is not None else "",
            str(last_updated) if last_updated is not None else ""
        ])

        workbook.save(EXCEL_FILE)
        print("User added to Excel successfully!")
    except Exception as e:
        print(f"Error writing to Excel file: {e}")
        raise
    finally:
        if workbook:
            workbook.close()

if __name__ == "__main__":
    create_users_excel()